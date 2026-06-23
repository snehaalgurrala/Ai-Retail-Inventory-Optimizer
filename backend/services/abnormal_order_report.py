"""Abnormal Order Intelligence Report — executive daily report (email + Excel).

Powers the dashboard "Abnormal Order Intelligence Reports" buttons. It reuses the
EXACT same abnormal-order pipeline as the Customer Intelligence page:

  * detection            -> ``customer_intelligence_service.detect_abnormal_orders``
  * card enrichment      -> ``abnormal_order_intelligence.build_abnormal_cards``
  * risk band / score    -> ``abnormal_order_intelligence._risk_assessment``
  * AI narrative & impact -> ``abnormal_order_intelligence`` narrative helpers

so risk classifications, priority scores, inventory scope and the AI investigation
wording all match the page. This module adds only the date filtering, the
executive email layout, and the Excel workbook — no new business logic.

The email is deliberately styled as an executive dashboard report (Bunzl blue +
green + white), not a simple alert. Sending reuses ``email_service.send_report_email``.
"""

from __future__ import annotations

from datetime import date, datetime, timedelta
from html import escape
from pathlib import Path

import pandas as pd

from backend.db import repository
from backend.services import abnormal_order_intelligence as aoi
from backend.services import customer_intelligence_service as cis
from backend.services.abnormal_order_email import _badge_html, _bullets, _metric_tile, _section
from backend.services.email_service import THEME, send_report_email


PROJECT_ROOT = Path(__file__).resolve().parents[2]
REPORTS_DIR = PROJECT_ROOT / "data" / "processed" / "reports"

_BAND_ACCENT = {
    "Critical": THEME["danger"],
    "High": THEME["warning"],
    "Medium": "#A86E00",
    "Low": THEME["fresh_green"],
}
_BAND_ORDER = ("Critical", "High", "Medium", "Low")


def _money(value: float) -> str:
    return f"${float(value):,.0f}"


# ---------------------------------------------------------------------------
# Data — same pipeline as the Customer Intelligence page, filtered to one day
# ---------------------------------------------------------------------------
def collect_abnormal_cards(target_date: date) -> list[dict]:
    """Return the enriched abnormal-order cards whose order_date == ``target_date``.

    Loads the live Oracle frames, runs the shared detection + card builder, then
    keeps only the lines placed on the requested day. Cards are ranked Critical →
    High → Medium → Low, then by composite score (descending), exactly like the page.
    """
    customers = repository.load_customers(safe=True)
    orders = repository.load_orders(safe=True)
    order_lines = repository.load_order_lines(safe=True)
    products = repository.load_products(safe=True)
    inventory = repository.load_inventory(safe=True)

    facts = cis.prepare_customer_orders(order_lines, orders, customers, products)
    abnormal_df = cis.detect_abnormal_orders(facts)
    if abnormal_df.empty or "order_date" not in abnormal_df.columns:
        return []

    target_str = target_date.strftime("%Y-%m-%d")
    day_df = abnormal_df[abnormal_df["order_date"].astype(str) == target_str].copy()
    if day_df.empty:
        return []

    at_risk_ids = cis.at_risk_products(inventory)
    cards = aoi.build_abnormal_cards(day_df, facts, at_risk_ids, inventory)

    cards.sort(
        key=lambda c: (
            aoi._RISK_PRIORITY[aoi._ensure_assessment(c)["band"]],
            -aoi._ensure_assessment(c)["score"],
        )
    )
    return cards


def _band_counts(cards: list[dict]) -> dict[str, int]:
    counts = {band: 0 for band in _BAND_ORDER}
    for card in cards:
        counts[aoi._ensure_assessment(card)["band"]] += 1
    return counts


# ---------------------------------------------------------------------------
# Email — executive dashboard layout
# ---------------------------------------------------------------------------
def _header_html(target_date: date, period_label: str, top_band: str) -> str:
    nice_date = target_date.strftime("%B %d, %Y")
    return (
        f'<div style="background:linear-gradient(135deg,{THEME["primary_navy"]},'
        f'{THEME["deep_navy"]});color:#FFFFFF;padding:34px 30px;text-align:center;">'
        '<div style="font-size:12px;font-weight:700;letter-spacing:1px;opacity:0.85;'
        'text-transform:uppercase;">Bunzl AI Supply Chain Platform · Executive Report</div>'
        '<h1 style="margin:8px 0 6px 0;font-size:26px;font-weight:800;">'
        '🚨 Abnormal Order Intelligence Report</h1>'
        f'<div style="font-size:13px;opacity:0.85;margin-bottom:14px;">'
        f'{escape(period_label)} · {escape(nice_date)}</div>'
        f'{_badge_html(top_band)}</div>'
    )


def _exec_summary_html(cards: list[dict], target_date: date, counts: dict[str, int]) -> str:
    total = len(cards)
    total_revenue = sum(float(c["revenue_impact"]) for c in cards)
    total_units = sum(int(c["current_quantity"]) for c in cards)
    nice_date = target_date.strftime("%b %d, %Y")

    def row(tiles: list[str]) -> str:
        return (
            '<table role="presentation" width="100%" cellpadding="0" cellspacing="0">'
            f'<tr>{"".join(tiles)}</tr></table>'
        )

    navy = THEME["primary_navy"]
    body = (
        row([
            _metric_tile("Report Date", nice_date, navy),
            _metric_tile("Total Abnormal Orders", f"{total:,}", THEME["danger"], highlight=True),
            _metric_tile("Critical Orders", f"{counts['Critical']:,}", _BAND_ACCENT["Critical"], highlight=True),
        ])
        + row([
            _metric_tile("High Risk Orders", f"{counts['High']:,}", _BAND_ACCENT["High"], highlight=True),
            _metric_tile("Medium Risk Orders", f"{counts['Medium']:,}", _BAND_ACCENT["Medium"], highlight=True),
            _metric_tile("Low Risk Orders", f"{counts['Low']:,}", _BAND_ACCENT["Low"], highlight=True),
        ])
        + row([
            _metric_tile("Total Revenue Impact", _money(total_revenue), THEME["fresh_green"], highlight=True),
            _metric_tile("Total Inventory Consumption", f"{total_units:,} units", navy),
            '<td style="width:33%;"></td>',
        ])
    )
    return _section("Executive Summary", body)


def _risk_distribution_html(counts: dict[str, int]) -> str:
    header = (
        f'<tr style="background:{THEME["primary_navy"]};color:#FFFFFF;">'
        '<th style="text-align:left;padding:10px 14px;font-size:13px;">Risk Level</th>'
        '<th style="text-align:right;padding:10px 14px;font-size:13px;">Count</th></tr>'
    )
    rows = []
    for band in _BAND_ORDER:
        accent = _BAND_ACCENT[band]
        rows.append(
            f'<tr style="border-bottom:1px solid {THEME["soft_border"]};">'
            f'<td style="padding:10px 14px;font-size:14px;font-weight:700;color:{accent};">'
            f'{escape(band)}</td>'
            f'<td style="padding:10px 14px;font-size:14px;font-weight:800;text-align:right;'
            f'color:{THEME["deep_navy"]};">{counts[band]:,}</td></tr>'
        )
    table = (
        '<table role="presentation" width="100%" cellpadding="0" cellspacing="0" '
        f'style="border-collapse:collapse;border:1px solid {THEME["soft_border"]};'
        'border-radius:10px;overflow:hidden;">'
        f'{header}{"".join(rows)}</table>'
    )
    return _section("Risk Distribution Summary", table)


_TABLE_COLUMNS = [
    "Order Date", "Order Number", "Customer Name", "Customer Segment", "Product Name",
    "Category", "Historical Average", "Historical Maximum", "Latest Order Quantity",
    "Deviation %", "Current Inventory", "Inventory Impact %", "Revenue Impact",
    "Risk Level", "Priority Score",
]


def _card_row_values(card: dict) -> dict[str, str]:
    ra = aoi._ensure_assessment(card)
    inv = card.get("current_inventory")
    impact = ra["impact_pct"]
    return {
        "Order Date": str(card.get("order_date") or "—"),
        "Order Number": str(card.get("order_nbr") or "—"),
        "Customer Name": str(card["customer_name"]),
        "Customer Segment": str(card.get("customer_segment") or "—"),
        "Product Name": str(card["product_name"]),
        "Category": str(card.get("category") or "—"),
        "Historical Average": f"{float(card['historical_avg']):,.1f}",
        "Historical Maximum": f"{int(card['hist_high']):,}",
        "Latest Order Quantity": f"{int(card['current_quantity']):,}",
        "Deviation %": f"+{float(card['deviation_pct']):.0f}%",
        "Current Inventory": f"{int(inv):,}" if inv is not None else "N/A",
        "Inventory Impact %": f"{impact:.0f}%" if impact is not None else "N/A",
        "Revenue Impact": _money(card["revenue_impact"]),
        "Risk Level": ra["band"],
        "Priority Score": f"{ra['score']}/100",
    }


def _dashboard_table_html(cards: list[dict]) -> str:
    header = "".join(
        '<th style="padding:9px 10px;text-align:left;font-size:11px;font-weight:700;'
        'text-transform:uppercase;letter-spacing:0.3px;white-space:nowrap;">'
        f'{escape(col)}</th>'
        for col in _TABLE_COLUMNS
    )
    rows = []
    for i, card in enumerate(cards):
        values = _card_row_values(card)
        band = values["Risk Level"]
        bg = "#FFFFFF" if i % 2 == 0 else THEME["soft_blue"]
        cells = []
        for col in _TABLE_COLUMNS:
            value = values[col]
            if col == "Risk Level":
                accent = _BAND_ACCENT.get(band, THEME["primary_navy"])
                cell = (
                    f'<span style="display:inline-block;padding:3px 10px;border-radius:999px;'
                    f'background:{accent};color:#FFFFFF;font-size:11px;font-weight:800;">'
                    f'{escape(band)}</span>'
                )
            else:
                cell = escape(value)
            cells.append(
                f'<td style="padding:9px 10px;font-size:12px;color:{THEME["deep_navy"]};'
                f'border-bottom:1px solid {THEME["soft_border"]};white-space:nowrap;">{cell}</td>'
            )
        rows.append(f'<tr style="background:{bg};">{"".join(cells)}</tr>')

    table = (
        '<div style="overflow-x:auto;border:1px solid ' + THEME["soft_border"] + ';'
        'border-radius:10px;">'
        '<table role="presentation" width="100%" cellpadding="0" cellspacing="0" '
        'style="border-collapse:collapse;min-width:1100px;">'
        f'<thead><tr style="background:{THEME["primary_navy"]};color:#FFFFFF;">{header}</tr></thead>'
        f'<tbody>{"".join(rows)}</tbody></table></div>'
    )
    return _section("Abnormal Orders Dashboard", table)


def _ai_investigation_html(cards: list[dict]) -> str:
    blocks = []
    for card in cards:
        ra = aoi._ensure_assessment(card)
        avg = max(1, round(float(card["historical_avg"])))
        cur = int(card["current_quantity"])
        dev = float(card["deviation_pct"])
        inv = card.get("current_inventory")
        impact = ra["impact_pct"]

        lines = [
            f"Historically this product averages {avg:,} units per order.",
            f"The latest order was {cur:,} units.",
            f"This represents a {dev:.0f}% increase compared with historical demand.",
        ]
        if impact is not None:
            lines.append(
                f"The order will consume approximately {impact:.0f}% of current network inventory."
            )
        elif inv is not None:
            lines.append(f"Current network inventory for this product is {int(inv):,} units.")

        investigation = "".join(
            f'<div style="font-size:14px;line-height:1.55;color:{THEME["deep_navy"]};'
            f'margin-bottom:4px;">{escape(line)}</div>'
            for line in lines
        )
        accent = _BAND_ACCENT.get(ra["band"], THEME["primary_navy"])
        block = (
            f'<div style="border:1px solid {THEME["soft_border"]};border-left:4px solid {accent};'
            'border-radius:10px;padding:14px 16px;margin-bottom:14px;background:#FFFFFF;">'
            f'<div style="font-size:15px;font-weight:800;color:{THEME["deep_navy"]};">'
            f'{escape(card["customer_name"])}</div>'
            f'<div style="font-size:13px;color:{THEME["muted_text"]};margin-bottom:10px;">'
            f'Product: {escape(card["product_name"])}</div>'
            f'<div style="font-size:11px;font-weight:700;text-transform:uppercase;letter-spacing:0.4px;'
            f'color:{accent};margin-bottom:6px;">AI Investigation</div>'
            f'{investigation}'
            f'<div style="margin-top:10px;font-size:13px;font-weight:800;color:{accent};">'
            f'Risk Level: {escape(ra["band"])}</div>'
            f'<div style="margin-top:10px;font-size:11px;font-weight:700;text-transform:uppercase;'
            f'letter-spacing:0.4px;color:{THEME["muted_text"]};">Potential Business Reasons</div>'
            f'{_bullets(aoi.BUSINESS_REASONS, THEME["primary_navy"])}'
            '</div>'
        )
        blocks.append(block)
    return _section("AI Investigation Summary", "".join(blocks))


def _business_impact_html(cards: list[dict]) -> str:
    blocks = []
    for card in cards:
        accent = _BAND_ACCENT.get(aoi._ensure_assessment(card)["band"], THEME["primary_navy"])
        block = (
            f'<div style="border:1px solid {THEME["soft_border"]};border-radius:10px;'
            f'padding:12px 16px;margin-bottom:12px;background:{THEME["soft_amber"]};">'
            f'<div style="font-size:14px;font-weight:800;color:{THEME["deep_navy"]};margin-bottom:6px;">'
            f'{escape(card["customer_name"])} · {escape(card["product_name"])}</div>'
            f'{_bullets(aoi.business_impact(card), accent)}'
            '</div>'
        )
        blocks.append(block)
    return _section("Business Impact", "".join(blocks))


def _recommended_actions_html(cards: list[dict]) -> str:
    blocks = []
    for card in cards:
        block = (
            f'<div style="border:1px solid {THEME["soft_border"]};border-left:4px solid '
            f'{THEME["fresh_green"]};border-radius:10px;padding:12px 16px;margin-bottom:12px;'
            f'background:{THEME["soft_blue"]};">'
            f'<div style="font-size:14px;font-weight:800;color:{THEME["deep_navy"]};margin-bottom:6px;">'
            f'{escape(card["customer_name"])} · {escape(card["product_name"])}</div>'
            f'{_bullets(aoi._deep_actions(card), THEME["fresh_green"])}'
            '</div>'
        )
        blocks.append(block)
    return _section("Recommended Actions", "".join(blocks))


def _footer_html() -> str:
    return (
        f'<div style="background:{THEME["deep_navy"]};color:#FFFFFF;padding:26px 30px;'
        'text-align:center;font-size:12px;line-height:1.7;">'
        '<div style="font-weight:700;opacity:0.95;">Generated by Bunzl AI Supply Chain '
        'Intelligence Platform</div>'
        '<div style="opacity:0.7;margin-top:6px;">Executive abnormal-order intelligence — '
        'figures match the Customer Intelligence dashboard.</div></div>'
    )


def build_report_html(cards: list[dict], target_date: date, period_label: str) -> str:
    counts = _band_counts(cards)
    top_band = next((band for band in _BAND_ORDER if counts[band] > 0), "Low")
    return (
        '<!DOCTYPE html><html lang="en"><head><meta charset="UTF-8">'
        '<meta name="viewport" content="width=device-width, initial-scale=1.0">'
        '<title>Abnormal Order Intelligence Report</title></head>'
        f'<body style="margin:0;padding:20px;background:{THEME["light_bg"]};'
        'font-family:-apple-system,Segoe UI,Helvetica,Arial,sans-serif;">'
        '<div style="max-width:1080px;margin:0 auto;background:#FFFFFF;border-radius:14px;'
        'overflow:hidden;box-shadow:0 10px 30px rgba(10,31,51,0.14);">'
        + _header_html(target_date, period_label, top_band)
        + _exec_summary_html(cards, target_date, counts)
        + _risk_distribution_html(counts)
        + _dashboard_table_html(cards)
        + _ai_investigation_html(cards)
        + _business_impact_html(cards)
        + _recommended_actions_html(cards)
        + _footer_html()
        + '</div></body></html>'
    )


# ---------------------------------------------------------------------------
# Excel attachment
# ---------------------------------------------------------------------------
_EXCEL_COLUMNS = _TABLE_COLUMNS + ["AI Recommendation", "Business Impact"]


def build_excel(cards: list[dict], target_date: date) -> Path:
    """Write the abnormal-order workbook and return its path."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Abnormal Orders"

    header_fill = PatternFill(start_color="183F5F", end_color="183F5F", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border = Border(
        left=Side(style="thin", color="D8E2EC"),
        right=Side(style="thin", color="D8E2EC"),
        top=Side(style="thin", color="D8E2EC"),
        bottom=Side(style="thin", color="D8E2EC"),
    )

    for col_num, title in enumerate(_EXCEL_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = title
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = border

    for row_num, card in enumerate(cards, 2):
        values = _card_row_values(card)
        for col_num, col in enumerate(_TABLE_COLUMNS, 1):
            ws.cell(row=row_num, column=col_num).value = values[col]
        ws.cell(row=row_num, column=len(_TABLE_COLUMNS) + 1).value = "; ".join(aoi._deep_actions(card))
        ws.cell(row=row_num, column=len(_TABLE_COLUMNS) + 2).value = " | ".join(aoi.business_impact(card))

    # Auto-size columns from content (capped so narrative columns stay readable).
    for col_num, title in enumerate(_EXCEL_COLUMNS, 1):
        letter = ws.cell(row=1, column=col_num).column_letter
        longest = len(title)
        for row_num in range(2, len(cards) + 2):
            text = str(ws.cell(row=row_num, column=col_num).value or "")
            longest = max(longest, len(text))
        ws.column_dimensions[letter].width = min(60, max(12, longest + 2))

    # Enable filters across the whole table.
    last_col = ws.cell(row=1, column=len(_EXCEL_COLUMNS)).column_letter
    ws.auto_filter.ref = f"A1:{last_col}{max(1, len(cards)) + 1}"
    ws.freeze_panes = "A2"

    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    output_path = REPORTS_DIR / f"Abnormal_Order_Report_{target_date.strftime('%Y%m%d')}.xlsx"
    wb.save(output_path)
    return output_path


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------
def _resolve_target(period: str, target_date: date | None) -> tuple[date, str, str]:
    """Return (date, period_label, subject_suffix) for the requested period."""
    today = datetime.now().date()
    if period == "today":
        return today, "Today's Report", "Today"
    if period == "yesterday":
        return today - timedelta(days=1), "Yesterday's Report", "Yesterday"
    chosen = target_date or today
    return chosen, f"Report for {chosen:%B %d, %Y}", chosen.strftime("%d %b %Y")


def send_abnormal_order_report_email(period: str = "today", target_date: date | None = None) -> dict:
    """Build and email the executive abnormal-order report for the chosen day.

    ``period`` is ``"today"``, ``"yesterday"`` or ``"date"`` (with ``target_date``).
    """
    resolved_date, period_label, subject_suffix = _resolve_target(period, target_date)

    try:
        cards = collect_abnormal_cards(resolved_date)
    except Exception as error:
        return {
            "success": False,
            "email_sent": False,
            "message": f"Could not load abnormal-order data: {error}",
        }

    if not cards:
        return {
            "success": False,
            "email_sent": False,
            "message": (
                f"No abnormal orders were detected for {resolved_date:%B %d, %Y}, "
                "so no report was sent."
            ),
        }

    html_body = build_report_html(cards, resolved_date, period_label)
    try:
        excel_path = build_excel(cards, resolved_date)
    except Exception as error:
        return {
            "success": False,
            "email_sent": False,
            "message": f"Could not build the Excel attachment: {error}",
        }

    subject = f"🚨 Bunzl Abnormal Order Intelligence Report – {subject_suffix}"
    result = send_report_email(subject=subject, html_body=html_body, attachment_path=excel_path)
    if result.get("success"):
        result["message"] = (
            f"Abnormal Order Intelligence Report ({len(cards)} order(s) for "
            f"{resolved_date:%B %d, %Y}) sent to the manager with "
            f"{excel_path.name} attached."
        )
    return result
